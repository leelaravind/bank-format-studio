"""P1-M5 tests: writers self-validate, reparse, round-trip, and are deterministic."""

from datetime import date, datetime
from decimal import Decimal

import pytest

from bfs_core.camt import V02, V08, read_camt053, write_camt053
from bfs_core.csvio import read_csv, write_csv
from bfs_core.model import (
    Balance,
    BankTransactionCode,
    Counterparty,
    CreditDebit,
    LossKind,
    Statement,
    Transaction,
)
from bfs_core.mt940 import read_mt940, write_mt940
from bfs_core.reconcile import reconcile_statement
from bfs_core.xlsx import write_xlsx

D = Decimal
C, DB = CreditDebit.CREDIT, CreditDebit.DEBIT
CLOCK = datetime(2026, 1, 15, 10, 30)


def rich_statement() -> Statement:
    return Statement(
        statement_id="SYNTH-W001",
        account_iban="DE75512108001245126199",
        account_currency="EUR",
        statement_number=7,
        opening_balance=Balance(C, date(2026, 1, 2), "EUR", D("1000.00")),
        closing_balance=Balance(C, date(2026, 1, 3), "EUR", D("1130.00")),
        closing_available=Balance(C, date(2026, 1, 3), "EUR", D("1100.00")),
        forward_available=(Balance(C, date(2026, 1, 4), "EUR", D("1130.00")),),
        transactions=[
            Transaction(
                value_date=date(2026, 1, 3), booking_date=date(2026, 1, 3),
                credit_debit=C, amount=D("200.00"), swift_tx_type="NTRF",
                customer_reference="CUST-1", bank_reference="BANK-1",
                end_to_end_id="E2E-0001", mandate_id="MND-9",
                counterparty=Counterparty(name="Acme Tooling GmbH",
                                          account="NL91ABNA0417164300", bic="ABNANL2A"),
                remittance_unstructured=("Invoice 4711",),
                purpose_code="GDDS",
            ),
            Transaction(
                value_date=date(2026, 1, 3), credit_debit=DB, amount=D("50.00"),
                is_reversal=True, swift_tx_type="NMSC",
                customer_reference="NONREF",
                btc=BankTransactionCode(domain="PMNT", family="RCDT", sub_family="OTHR"),
            ),
            Transaction(
                value_date=date(2026, 1, 3), credit_debit=DB, amount=D("20.00"),
                swift_tx_type="NCHG", funds_code="R",
                supplementary_details="fee note",
                charges_amount=D("20.00"),
            ),
        ],
        source_format="test",
    )


class TestCamtWriters:
    @pytest.mark.parametrize("spec", [V02, V08], ids=["v02", "v08"])
    def test_output_validates_and_reparses(self, spec):
        s = rich_statement()
        payload, report = write_camt053([s], spec, CLOCK)
        statements, read_report = read_camt053(payload)
        assert not read_report.has_errors
        s2 = statements[0]
        r1, r2 = reconcile_statement(s), reconcile_statement(s2)
        assert r1.passed and r2.passed
        assert r1.conservation_key() == r2.conservation_key()  # INV-7
        t = s2.transactions[0]
        assert t.end_to_end_id == "E2E-0001"
        assert t.counterparty and t.counterparty.name == "Acme Tooling GmbH"
        assert s2.transactions[1].is_reversal  # RvslInd round-trip

    @pytest.mark.parametrize("spec", [V02, V08], ids=["v02", "v08"])
    def test_deterministic(self, spec):
        a, _ = write_camt053([rich_statement()], spec, CLOCK)
        b, _ = write_camt053([rich_statement()], spec, CLOCK)
        assert a == b

    def test_loss_notes_for_mt940_only_fields(self):
        _, report = write_camt053([rich_statement()], V02, CLOCK)
        fields = {n.field_name for n in report.loss_notes}
        assert "funds_code" in fields
        assert "supplementary_details" in fields

    def test_batch_entry_v02_and_v08(self):
        s = rich_statement()
        s.transactions = [Transaction(
            value_date=date(2026, 1, 3), credit_debit=C, amount=D("130.00"),
            swift_tx_type="NTRF", entry_reference="BATCH-1",
            details=(
                Transaction(value_date=date(2026, 1, 3), credit_debit=C,
                            amount=D("100.00"), end_to_end_id="E2E-A"),
                Transaction(value_date=date(2026, 1, 3), credit_debit=C,
                            amount=D("30.00"), end_to_end_id="E2E-B"),
            ),
        )]
        for spec in (V02, V08):
            payload, _ = write_camt053([s], spec, CLOCK)
            statements, report = read_camt053(payload)
            assert not report.has_errors
            entry = statements[0].transactions[0]
            assert len(entry.details) == 2
            assert sum(d.amount for d in entry.details) == D("130.00")


class TestMt940Writer:
    def test_reparses_and_reconciles(self):
        s = rich_statement()
        payload, report = write_mt940([s])
        statements, read_report = read_mt940(payload)
        assert len(statements) == 1
        s2 = statements[0]
        r1, r2 = reconcile_statement(s), reconcile_statement(s2)
        assert r2.passed
        assert r1.conservation_key() == r2.conservation_key()  # INV-7

    def test_gate4_slash_convention_round_trip(self):
        s = rich_statement()
        payload, _ = write_mt940([s])
        text = payload.decode("ascii")
        assert "/EREF/E2E-0001" in text
        assert "/ORDP//NAME/Acme Tooling GmbH" in text
        assert "/REMI/Invoice 4711" in text
        assert "/PURP//CD/GDDS" in text
        statements, _ = read_mt940(payload)
        t = statements[0].transactions[0]
        assert t.end_to_end_id == "E2E-0001"
        assert t.counterparty and t.counterparty.account == "NL91ABNA0417164300"
        assert t.counterparty.bic == "ABNANL2A"

    def test_reversal_marks_written(self):
        payload, _ = write_mt940([rich_statement()])
        assert b"RC50,00" in payload  # DEBIT reversal → RC (GATE-1 inverse)

    def test_long_reference_truncated_with_loss_note(self):
        # B-1 remediation: :61: still truncates to 16x, but the FULL value is
        # spilled to :86:/CREF/ and restored on reparse (previously it was lost
        # and the loss note falsely claimed preservation).
        s = rich_statement()
        s.transactions[0].customer_reference = "X" * 20
        payload, report = write_mt940([s])
        assert any(n.kind is LossKind.TRUNCATED for n in report.loss_notes)
        assert b"/CREF/" + b"X" * 20 in payload.replace(b"\r\n", b"")
        statements, _ = read_mt940(payload)
        assert statements[0].transactions[0].customer_reference == "X" * 20

    def test_transliteration_recorded(self):
        s = rich_statement()
        s.transactions[0].remittance_unstructured = ("Überweisung März",)
        payload, report = write_mt940([s])
        assert b"Uberweisung Marz" in payload
        assert any(n.kind is LossKind.TRANSLITERATED for n in report.loss_notes)

    def test_deterministic(self):
        assert write_mt940([rich_statement()])[0] == write_mt940([rich_statement()])[0]


class TestCsv:
    def test_round_trip_conserves(self):
        s = rich_statement()
        tx_csv, st_csv, report = write_csv([s])
        statements, _ = read_csv(tx_csv, st_csv)
        s2 = statements[0]
        r1, r2 = reconcile_statement(s), reconcile_statement(s2)
        assert r2.passed
        assert r1.conservation_key() == r2.conservation_key()
        assert s2.transactions[1].is_reversal
        assert s2.transactions[0].customer_reference == "CUST-1"

    def test_mandatory_statements_csv(self):
        tx_csv, st_csv, _ = write_csv([rich_statement()])
        assert st_csv.startswith(b"\xef\xbb\xbf")
        assert b"statement_id" in st_csv and b"total_credits" in st_csv

    def test_formula_neutralization_default_on(self):
        s = rich_statement()
        s.transactions[0].remittance_unstructured = ("=cmd|calc", )
        tx_csv, _, _ = write_csv([s])
        assert b"'=cmd|calc" in tx_csv
        tx_csv_strict, _, _ = write_csv([s], excel_safe=False)
        assert b"'=cmd|calc" not in tx_csv_strict

    def test_neutralization_reversed_on_import(self):
        s = rich_statement()
        s.transactions[0].remittance_unstructured = ("=cmd|calc",)
        tx_csv, st_csv, _ = write_csv([s])
        statements, _ = read_csv(tx_csv, st_csv)
        assert statements[0].transactions[0].remittance_unstructured == ("=cmd|calc",)

    def test_signed_amounts_dot_decimal_crlf_bom(self):
        tx_csv, _, _ = write_csv([rich_statement()])
        text = tx_csv.decode("utf-8-sig")
        assert "\r\n" in text
        assert "-50.00" in text and "200.00" in text

    def test_deterministic(self):
        assert write_csv([rich_statement()])[0] == write_csv([rich_statement()])[0]


class TestXlsx:
    def test_two_sheets_with_native_types(self):
        payload, _ = write_xlsx([rich_statement()])
        from io import BytesIO

        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(payload))
        assert wb.sheetnames == ["transactions", "statements"]
        tx = wb["transactions"]
        header = [c.value for c in tx[1]]
        amount_col = header.index("amount") + 1
        value = tx.cell(row=2, column=amount_col).value
        assert isinstance(value, (int, float, Decimal))

    def test_formula_cells_neutralized(self):
        s = rich_statement()
        s.transactions[0].remittance_unstructured = ("@SUM(A1)",)
        payload, _ = write_xlsx([s])
        from io import BytesIO

        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(payload))
        tx = wb["transactions"]
        header = [c.value for c in tx[1]]
        col = header.index("remittance_info") + 1
        assert tx.cell(row=2, column=col).value == "'@SUM(A1)"
